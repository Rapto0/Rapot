"""Run the real synthetic CLI in an isolated process; inspect actual report artifacts."""

import builtins
import hashlib
import importlib
import json
import os
import subprocess
import sys
from pathlib import Path
from types import ModuleType
from typing import Any
from xml.etree import ElementTree

import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def cli_runs(tmp_path_factory: pytest.TempPathFactory) -> tuple[Path, Path]:
    root = tmp_path_factory.mktemp("full-backtest-cli")
    working = root / "working"
    working.mkdir()
    # Reading caller dotenv or trusting caller application settings must not be necessary.
    (working / ".env").write_text("SCAN_INTERVAL_HOURS=not-a-number\n", encoding="utf-8")
    environment = dict(os.environ)
    environment.update({"PYTHONPATH": str(ROOT), "SCAN_INTERVAL_HOURS": "invalid-caller-setting"})
    outputs = (root / "first", root / "second")
    for output in outputs:
        result = subprocess.run(
            [
                sys.executable,
                "-X",
                "utf8",
                "-B",
                "-m",
                "scripts.backtest_fixture",
                "--output-dir",
                str(output),
                "--as-of",
                "2020-06-16T00:00:00Z",
                "--excel",
            ],
            cwd=working,
            env=environment,
            capture_output=True,
            text=True,
            encoding="utf-8",
            timeout=180,
            check=False,
        )
        assert result.returncode == 0, result.stdout + result.stderr
        assert "Report:" in result.stdout
        assert "prohibited IO" not in result.stderr
    assert sorted(path.name for path in working.iterdir()) == [".env"]
    return outputs


def test_real_fixture_cli_is_reproducible_and_creates_only_declared_artifacts(
    cli_runs: tuple[Path, Path],
) -> None:
    first, second = cli_runs
    expected = {
        "report.json",
        "report.xlsx",
        "equity.svg",
        *(
            f"{market}_{kind}.csv"
            for market in ("bist", "crypto")
            for kind in ("trades", "equity", "open_positions")
        ),
    }
    assert {path.name for path in first.iterdir()} == expected
    assert {path.name for path in second.iterdir()} == expected
    # XLSX ZIP/container timestamps are presentation metadata, not numerical reproducibility.
    for filename in expected - {"report.xlsx"}:
        assert (first / filename).read_bytes() == (second / filename).read_bytes()
    report = json.loads((first / "report.json").read_text(encoding="utf-8"))
    assert report["mode"] == "synthetic_fixture"
    assert report["fixture_id"] == "daily-waves-v1"
    assert report["as_of"] == "2020-06-16T00:00:00+00:00"
    assert len(report["input_sha256"]) == 4
    assert report["walk_forward_optimization_performed"] is False
    for market in ("BIST", "CRYPTO"):
        entry = report["markets"][market]
        assert entry["trade_count"] > 0, (
            "Real calculators must exercise accounting, not only empty output"
        )
        assert entry["equity_count"] == 27
        stats = entry["statistics"]
        assert stats["net_asset_value"] is not None
        assert stats["open_lots"] > 0
        assert stats["realized_profit"] == 0
        assert stats["total_profit"] == pytest.approx(stats["unrealized_profit"])
        assert report["benchmark"][market]["benchmark_return"] is not None
        assert report["rolling_buy_and_hold"][market]["strategy_evaluated"] is False
    for filename, manifest in report["csv_artifacts"].items():
        data = (first / filename).read_bytes()
        assert manifest == {"bytes": len(data), "sha256": hashlib.sha256(data).hexdigest()}


def test_actual_excel_values_and_svg_points_match_exported_nav(cli_runs: tuple[Path, Path]) -> None:
    from openpyxl import load_workbook

    output = cli_runs[0]
    report = json.loads((output / "report.json").read_text(encoding="utf-8"))
    workbook = load_workbook(output / "report.xlsx", read_only=True, data_only=True)
    try:
        summary = {
            row[0]: row[1:] for row in workbook["Genel Özet"].iter_rows(min_row=2, values_only=True)
        }
        assert "Getiri %" not in summary
        for column, market in enumerate(("BIST", "CRYPTO")):
            stats = report["markets"][market]["statistics"]
            currency = "TL" if market == "BIST" else "USD"
            assert (
                summary["Son Net Varlık Değeri (NAV)"][column]
                == f"{stats['net_asset_value']:,.2f} {currency}"
            )
            assert summary["Gerçekleşen Kar/Zarar"][column] == "0.00 " + currency
            assert summary["NAV Getiri %"][column] == f"{stats['total_return_pct']:.2f}%"
            equity = pd.read_csv(output / f"{market.lower()}_equity.csv")
            assert equity.iloc[-1]["Toplam Değer"] == stats["net_asset_value"]
    finally:
        workbook.close()
    svg = ElementTree.parse(output / "equity.svg").getroot()
    lines = svg.findall(".//{http://www.w3.org/2000/svg}polyline")
    assert len(lines) == 2
    assert all(len(line.attrib["points"].split()) == 27 for line in lines)
    assert "closing NAV" in svg.find("{http://www.w3.org/2000/svg}title").text


def test_cli_refuses_to_overwrite_existing_report(cli_runs: tuple[Path, Path]) -> None:
    output = cli_runs[0]
    before = (output / "report.json").read_bytes()
    result = subprocess.run(
        [
            sys.executable,
            "-X",
            "utf8",
            "-B",
            "-m",
            "scripts.backtest_fixture",
            "--output-dir",
            str(output),
        ],
        cwd=ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=30,
        check=False,
    )
    assert result.returncode != 0
    assert "never overwritten" in result.stderr
    assert (output / "report.json").read_bytes() == before


def test_fixture_module_import_does_not_load_application_or_create_files(tmp_path: Path) -> None:
    working = tmp_path / "import-only"
    working.mkdir()
    environment = {**os.environ, "PYTHONPATH": str(ROOT)}
    result = subprocess.run(
        [
            sys.executable,
            "-B",
            "-c",
            "import sys; import scripts.backtest_fixture; assert 'settings' not in sys.modules; assert 'backtesting_system' not in sys.modules",
        ],
        cwd=working,
        env=environment,
        capture_output=True,
        text=True,
        timeout=30,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    assert list(working.iterdir()) == []


def test_fixture_guard_rejects_network_database_and_provider_attempts(tmp_path: Path) -> None:
    # Exercise audit events directly: a broken guard cannot contact a host or open a DB.
    code = """
import sys
from scripts.backtest_fixture import main
try:
    main(['--output-dir', sys.argv[1], '--as-of', 'not-a-timestamp'])
except ValueError:
    pass
else:
    raise AssertionError('Invalid cutoff was accepted')
for event in ('socket.connect', 'socket.getaddrinfo', 'sqlite3.connect', 'subprocess.Popen'):
    try:
        sys.audit(event, None)
    except RuntimeError as error:
        assert 'prohibited IO' in str(error)
    else:
        raise AssertionError(event + ' was not blocked')
try:
    sys.modules['backtesting_system'].get_crypto_data('FAKE')
except RuntimeError as error:
    assert 'provider access' in str(error)
else:
    raise AssertionError('Provider was not blocked')
"""
    result = subprocess.run(
        [sys.executable, "-X", "utf8", "-B", "-c", code, str(tmp_path / "output")],
        cwd=tmp_path,
        env={**os.environ, "PYTHONPATH": str(ROOT)},
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=30,
        check=False,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert not (tmp_path / "output").exists()


@pytest.fixture
def backtest() -> ModuleType:
    return importlib.import_module("backtesting_system")


def test_stats_distinguish_unrealized_nav_from_realized_profit(backtest: ModuleType) -> None:
    costs = backtest.TradingCosts(bist_commission=0, bist_slippage=0)
    portfolio = backtest.Portfolio(200, "BIST", 100, costs)
    engine = backtest.BacktestEngine()
    assert portfolio.buy("AAA", 100, pd.Timestamp("2020-01-01"), "test")
    unmarked = engine._calculate_stats(portfolio)
    assert unmarked["net_asset_value"] is None
    assert unmarked["total_return_pct"] is None
    assert unmarked["open_cost_basis"] == 100
    portfolio.record_equity(pd.Timestamp("2020-01-02"), {"AAA": 150})
    marked = engine._calculate_stats(portfolio)
    assert marked["profit"] == marked["realized_profit"] == 0
    assert marked["net_asset_value"] == 250
    assert marked["unrealized_profit"] == marked["total_profit"] == 50
    assert marked["total_return_pct"] == 25
    assert portfolio.buy("AAA", 100, pd.Timestamp("2020-01-03"), "later")
    assert engine._calculate_stats(portfolio)["net_asset_value"] is None


def test_same_day_same_cash_replacement_cannot_reuse_an_old_mark(backtest: ModuleType) -> None:
    portfolio = backtest.Portfolio(
        100, "BIST", 100, backtest.TradingCosts(bist_commission=0, bist_slippage=0)
    )
    engine = backtest.BacktestEngine()
    day = pd.Timestamp("2020-01-01")
    assert portfolio.buy("AAA", 100, day, "original")
    portfolio.record_equity(day, {"AAA": 150})
    assert engine._calculate_stats(portfolio)["net_asset_value"] == 150
    assert portfolio.sell("AAA", 100, day, "replace")
    assert portfolio.buy("BBB", 100, day, "replacement")
    assert portfolio.cash == 0
    assert engine._calculate_stats(portfolio)["net_asset_value"] is None


def test_empty_portfolio_reports_are_valid(backtest: ModuleType, tmp_path: Path) -> None:
    from scripts.backtest_fixture import write_reports

    engine = backtest.BacktestEngine(as_of="2020-06-16T00:00:00Z")
    portfolios = {market: backtest.Portfolio(100, market, 10) for market in ("BIST", "CRYPTO")}
    stats = engine._calculate_stats(portfolios["BIST"])
    assert stats["net_asset_value"] == 100 and stats["total_return_pct"] == 0
    engine.plot_results(*portfolios.values(), output_path=tmp_path / "empty.svg")
    assert (
        not ElementTree.parse(tmp_path / "empty.svg")
        .getroot()
        .findall(".//{http://www.w3.org/2000/svg}polyline")
    )
    path = write_reports(engine, portfolios, tmp_path, {"mode": "test"})
    assert json.loads(path.read_text(encoding="utf-8"))["markets"]["BIST"]["trade_count"] == 0
    assert pd.read_csv(tmp_path / "bist_trades.csv").empty


def test_missing_optional_renderers_have_actionable_errors(
    backtest: ModuleType,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    engine = backtest.BacktestEngine()
    portfolio = backtest.Portfolio(100, "BIST", 10)
    original = builtins.__import__

    def without_matplotlib(name: str, *args: Any, **kwargs: Any) -> Any:
        if name.startswith("matplotlib"):
            raise ImportError("Synthetic absent optional renderer")
        return original(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", without_matplotlib)
    with pytest.raises(RuntimeError, match="use .svg"):
        engine.plot_results(portfolio, portfolio, output_path=tmp_path / "absent.png")

    def missing_excel(*args: Any, **kwargs: Any) -> None:
        raise ImportError("Synthetic absent optional workbook dependency")

    monkeypatch.setattr(backtest.pd, "ExcelWriter", missing_excel)
    with pytest.raises(RuntimeError, match="openpyxl"):
        engine.generate_excel_report(portfolio, portfolio, output_path=tmp_path / "absent.xlsx")
